"""Evaluate every formula with pycel, then check key outputs against an independent Python model."""
import math, sys
from openpyxl import load_workbook
from pycel import ExcelCompiler

path = sys.argv[1]
wb = load_workbook(path)
xl = ExcelCompiler(filename=path)

errors, count = [], 0
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                count += 1
                v = xl.evaluate("%s!%s" % (ws.title, c.coordinate))
                if isinstance(v, str) and v.startswith("#"):
                    errors.append((ws.title, c.coordinate, v, c.value))
print("formulas evaluated:", count, "| errors:", len(errors))
for e in errors[:10]: print("  ", e)

def model(price, tickers, days, llm, laya, fee_pct, fee_fixed, start, new1, g, churn, cac, hosting, neat, team, cash0):
    cust, cash, out = start, cash0, {"cash": [], "op": [], "cust": [], "rev": []}
    for m in range(1, 25):
        new = new1 * (1 + g) ** (m - 1)
        end = cust + new - cust * churn
        rev = end * price
        gross = rev - end * (price * fee_pct + fee_fixed) - end * tickers * days * 0.123 * (llm + laya)
        op = gross - new * cac - hosting - neat - team
        cash += op
        for k, v in (("cash", cash), ("op", op), ("cust", end), ("rev", rev)): out[k].append(v)
        cust = end
    return out

S = {"Base": model(19, 5, 21, 0.038, 0, 0.029, 0.30, 0, 10, 0.05, 0.08, 30, 100, 0, 2000, 25000),
     "Growth": model(19, 5, 21, 0.038, 0, 0.029, 0.30, 0, 25, 0.10, 0.06, 40, 100, 0, 6000, 25000)}
checks = []
for col, name in (("B", "Base"), ("C", "Growth")):
    o = S[name]
    first_neg = next((i + 1 for i, v in enumerate(o["cash"]) if v < 0), "none within 24 months")
    first_profit = next((i + 1 for i, v in enumerate(o["op"]) if v >= 0), "not within 24 months")
    expect = {5: o["cust"][11], 6: o["cust"][23], 7: o["rev"][11], 8: o["rev"][23], 11: sum(o["op"]),
              12: min(o["cash"]), 13: o["cash"][23], 14: first_neg, 15: first_profit, 16: max(0, -min(o["cash"]))}
    for row, want in expect.items():
        got = xl.evaluate("Summary!%s%d" % (col, row))
        ok = (got == want) if isinstance(want, str) else math.isclose(float(got), want, rel_tol=1e-9, abs_tol=1e-6)
        checks.append(ok)
        print("%-7s Summary!%s%-3d %-28s sheet=%s python=%s %s" % (name, col, row, wb["Summary"].cell(row=row, column=1).value[:28],
              round(got, 2) if not isinstance(got, str) else got, round(want, 2) if not isinstance(want, str) else want, "OK" if ok else "MISMATCH"))
cm = 19 - (19 * 0.029 + 0.30) - 105 * 0.123 * 0.038
for col, churn, cac in (("B", 0.08, 30), ("C", 0.06, 40)):
    for row, want in ((8, cm), (12, cm / churn), (13, cm / churn / cac)):
        got = xl.evaluate("Unit_Economics!%s%d" % (col, row))
        ok = math.isclose(got, want, rel_tol=1e-9); checks.append(ok)
        print("Unit_Economics!%s%d sheet=%.3f python=%.3f %s" % (col, row, got, want, "OK" if ok else "MISMATCH"))
print("ALL CHECKS PASS" if all(checks) and not errors else "CHECKS FAILED")
