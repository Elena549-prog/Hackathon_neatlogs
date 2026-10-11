"""Build results/business_model.xlsx: the trading-swarm business model for import into cfo.ai.

Every number the model uses is an input on the Assumptions sheet; every other cell is a formula.
"""
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = sys.argv[1]
MONTHS = 24
FIRST_COL = 3                      # month 1 in column C
LAST = get_column_letter(FIRST_COL + MONTHS - 1)

FONT = "Arial"
BLUE, BLACK, GREEN = "0000FF", "000000", "008000"
YELLOW = PatternFill("solid", fgColor="FFFF00")
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
THIN = Border(bottom=Side(style="thin", color="999999"))
USD = '$#,##0;($#,##0);"-"'
USD_CENTS = '$#,##0.00;($#,##0.00);"-"'
USD_MILLS = '$0.000;($0.000);"-"'
PCT = '0.0%;(0.0%);"-"'
COUNT = '#,##0.0;(#,##0.0);"-"'
INT = '#,##0;(#,##0);"-"'


def font(color=BLACK, bold=False, size=10, italic=False):
    return Font(name=FONT, color=color, bold=bold, size=size, italic=italic)


def style_all(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font.name != FONT:
                c.font = Font(name=FONT, size=c.font.size or 10, bold=c.font.bold,
                              italic=c.font.italic, color=c.font.color)


wb = Workbook()

# ------------------------------------------------------------------ Assumptions
A = wb.active
A.title = "Assumptions"
A["A1"] = "Trading swarm: business model assumptions"
A["A1"].font = font(bold=True, size=13)
A["A2"] = ("Legend: blue = input you can change. Yellow = key assumption to confirm (owner's estimate, "
           "not measured). Black = formula. Green = link to another sheet. Column B feeds Model_Base, "
           "column C feeds Model_Growth.")
A["A2"].font = font(italic=True, size=9)
A["A3"] = ("Product: a pre-trade gate for paper and retail algo traders. For each watchlist ticker, every "
           "trading day, a four-agent swarm decides whether a +5%/-3%/10-day trade has evidence behind it, "
           "and refuses when it doesn't. Priced per watchlist (5 tickers) per month.")
A["A3"].font = font(size=9)
A["A3"].alignment = Alignment(wrap_text=True, vertical="top")
A.merge_cells("A3:E3")
A.row_dimensions[3].height = 40

headers = ["Assumption", "Base", "Growth", "Unit", "Source or note"]
for i, h in enumerate(headers, 1):
    c = A.cell(row=5, column=i, value=h)
    c.font = font(bold=True)
    c.fill = HEADER_FILL
    c.border = THIN

# key, label, base, growth, unit, number format, key assumption?, source
ROWS = [
    ("price", "Price per watchlist per month", 19, 19, "$ / month", USD_CENTS, True,
     "Owner assumption: a hobbyist-tier subscription price. Not tested with customers."),
    ("tickers", "Tickers per watchlist", 5, 5, "tickers", INT, False,
     "Matches the project's 5-ticker watchlist (AAPL, MSFT, GOOGL, AMZN, NVDA)."),
    ("days", "Trading days per month", 21, 21, "days", INT, False,
     "About 252 US trading days a year / 12."),
    ("trigger", "Share of days the price watcher wakes the swarm", 0.123, 0.123, "%", PCT, False,
     "Measured on the train split: a 1-day move of at least 1.5x the 20-day volatility on 12.3% of days "
     "(17.0% on validation). Quiet days cost nothing: no LLM call (docs/SWARM_SPEC.md section 15)."),
    ("decisions", "Swarm decisions per watchlist per month", "=B7*B8*B9", "=C7*C8*C9", "decisions", COUNT, False,
     "Formula: tickers x trading days x watcher share."),
    ("llm", "LLM cost per decision", 0.038, 0.038, "$ / decision", USD_MILLS, False,
     "Measured: $11.76 logged over 310 validation setups on 2026-10-10 (results/README.md), with two Planner "
     "rounds per setup. Upper bound: DeepSeek peak, cache-miss prices; off-peak is half. The second round now "
     "only runs for fixable failures, so the real figure should be lower; re-measure with the next run."),
    ("laya", "Laya inference cost per decision", 0, 0, "$ / decision", USD_MILLS, False,
     "Laya is a local classifier (about 33 ms per question on a T4 GPU). Marginal cost about zero; "
     "set above zero if it runs on paid GPUs. Not used in production unless the step 9 go/no-go passes."),
    ("fee_pct", "Payment fee, share of price", 0.029, 0.029, "%", PCT, False,
     "Stripe standard US card pricing: 2.9% + $0.30 per successful charge."),
    ("fee_fixed", "Payment fee, fixed per charge", 0.30, 0.30, "$ / charge", USD_CENTS, False,
     "Stripe standard US card pricing."),
    ("start_customers", "Customers at start", 0, 0, "customers", INT, False, "Pre-launch."),
    ("new1", "New customers in month 1", 10, 25, "customers", INT, True,
     "Owner assumption. Base: organic from build-in-public posts. Growth: plus paid acquisition."),
    ("new_growth", "Monthly growth in new customers", 0.05, 0.10, "% / month", PCT, True,
     "Owner assumption."),
    ("churn", "Monthly churn", 0.08, 0.06, "% / month", PCT, True,
     "Owner assumption, typical range for consumer subscriptions."),
    ("cac", "Acquisition cost per new customer (CAC)", 30, 40, "$ / customer", USD, True,
     "Owner assumption. Growth spends more per customer on paid channels."),
    ("hosting", "Hosting and tools per month", 100, 100, "$ / month", USD, True,
     "Owner assumption: a small cloud VM plus a database."),
    ("neatlogs", "neatlogs per month", 0, 0, "$ / month", USD, True,
     "Placeholder: hackathon credits. Replace with the neatlogs plan price."),
    ("team", "Team cost per month", 2000, 6000, "$ / month", USD, True,
     "Owner assumption. Base: two part-time founders' stipend. Growth: adds one hire."),
    ("cash", "Starting cash", 25000, 25000, "$", USD, True,
     "Owner assumption: founder-funded."),
]
REF = {}            # key -> row on Assumptions
for i, (key, label, base, growth, unit, fmt, key_assumption, note) in enumerate(ROWS):
    r = 6 + i
    REF[key] = r
    A.cell(row=r, column=1, value=label).font = font()
    for col, val in ((2, base), (3, growth)):
        c = A.cell(row=r, column=col, value=val)
        c.number_format = fmt
        is_formula = isinstance(val, str) and val.startswith("=")
        c.font = font(BLACK if is_formula else BLUE)
        if key_assumption and not is_formula:
            c.fill = YELLOW
    A.cell(row=r, column=4, value=unit).font = font(size=9)
    n = A.cell(row=r, column=5, value=note)
    n.font = font(size=9)
    n.alignment = Alignment(wrap_text=True, vertical="top")
    A.row_dimensions[r].height = 30 if len(note) > 70 else 15
A["B11"].comment = Comment("Measured upper bound from the 2026-10-10 validation run. See results/README.md.", "trading-swarm")
for col, width in zip("ABCDE", (38, 12, 12, 13, 80)):
    A.column_dimensions[col].width = width
A.freeze_panes = "A6"


def a(key, col):    # absolute reference to an assumption for a scenario column (B or C)
    return "Assumptions!$%s$%d" % (col, REF[key])


# --------------------------------------------------------------- monthly models
MODEL_ROWS = {}     # label key -> row number on model sheets


def build_model(title, col):
    ws = wb.create_sheet(title)
    ws["A1"] = "%s scenario: monthly model (months 1-%d)" % (title.split("_")[1], MONTHS)
    ws["A1"].font = font(bold=True, size=13)
    ws["A2"] = "All values in $ unless the unit column says otherwise. Inputs live on the Assumptions sheet, column %s." % col
    ws["A2"].font = font(italic=True, size=9)

    lines = [   # key, label, unit, first-month formula, later-month formula ({c} = column, {p} = previous column), format, color
        ("month", "Month", "", "1", "={p}{r_month}+1", INT, BLACK),
        None,
        ("cust_start", "Customers, start of month", "customers", "=" + a("start_customers", col), "={p}{r_cust_end}", COUNT, None),
        ("new", "New customers", "customers",
         "=%s*(1+%s)^({c}{r_month}-1)" % (a("new1", col), a("new_growth", col)), None, COUNT, BLACK),
        ("churned", "Churned customers", "customers", "={c}{r_cust_start}*%s" % a("churn", col), None, COUNT, BLACK),
        ("cust_end", "Customers, end of month", "customers", "={c}{r_cust_start}+{c}{r_new}-{c}{r_churned}", None, COUNT, BLACK),
        None,
        ("revenue", "Revenue (MRR)", "$", "={c}{r_cust_end}*%s" % a("price", col), None, USD, BLACK),
        ("fees", "Payment fees", "$", "={c}{r_cust_end}*(%s*%s+%s)" % (a("price", col), a("fee_pct", col), a("fee_fixed", col)), None, USD, BLACK),
        ("llm", "LLM inference", "$", "={c}{r_cust_end}*%s*%s" % (a("decisions", col), a("llm", col)), None, USD, BLACK),
        ("laya", "Laya inference", "$", "={c}{r_cust_end}*%s*%s" % (a("decisions", col), a("laya", col)), None, USD, BLACK),
        ("gross", "Gross profit", "$", "={c}{r_revenue}-{c}{r_fees}-{c}{r_llm}-{c}{r_laya}", None, USD, BLACK),
        ("gm", "Gross margin", "%", "=IF({c}{r_revenue}=0,0,{c}{r_gross}/{c}{r_revenue})", None, PCT, BLACK),
        None,
        ("marketing", "Marketing (new customers x CAC)", "$", "={c}{r_new}*%s" % a("cac", col), None, USD, BLACK),
        ("hosting", "Hosting and tools", "$", "=" + a("hosting", col), None, USD, GREEN),
        ("neatlogs", "neatlogs", "$", "=" + a("neatlogs", col), None, USD, GREEN),
        ("team", "Team", "$", "=" + a("team", col), None, USD, GREEN),
        ("op", "Operating profit", "$", "={c}{r_gross}-{c}{r_marketing}-{c}{r_hosting}-{c}{r_neatlogs}-{c}{r_team}", None, USD, BLACK),
        None,
        ("cash_start", "Cash, start of month", "$", "=" + a("cash", col), "={p}{r_cash_end}", USD, None),
        ("cash_end", "Cash, end of month", "$", "={c}{r_cash_start}+{c}{r_op}", None, USD, BLACK),
        None,
        ("flag_cash", "Flag: cash below zero (1 = yes)", "flag", "=IF({c}{r_cash_end}<0,1,0)", None, INT, BLACK),
        ("flag_profit", "Flag: operating profit at or above zero (1 = yes)", "flag", "=IF({c}{r_op}>=0,1,0)", None, INT, BLACK),
    ]
    r = 4
    rows = {}
    for line in lines:
        if line:
            rows[line[0]] = r
        r += 1
    MODEL_ROWS.update(rows)
    names = {"r_" + k: v for k, v in rows.items()}

    ws.cell(row=3, column=1, value="Line").font = font(bold=True)
    ws.cell(row=3, column=2, value="Unit").font = font(bold=True)
    for line in lines:
        if not line:
            continue
        key, label, unit, first, later, fmt, color = line
        rr = rows[key]
        ws.cell(row=rr, column=1, value=label).font = font(bold=key in ("cust_end", "revenue", "gross", "op", "cash_end"))
        ws.cell(row=rr, column=2, value=unit).font = font(size=9)
        for m in range(MONTHS):
            colno = FIRST_COL + m
            c_letter = get_column_letter(colno)
            p_letter = get_column_letter(colno - 1)
            tmpl = first if (m == 0 or later is None) else later
            value = tmpl.format(c=c_letter, p=p_letter, **names)
            cell = ws.cell(row=rr, column=colno, value=int(value) if value.isdigit() else value)
            cell.number_format = fmt
            if key == "month":
                cell.font = font(BLUE if m == 0 else BLACK, bold=True)
                cell.fill = HEADER_FILL
            else:
                link = value.startswith("=Assumptions!") or (color is None and m == 0)
                cell.font = font(GREEN if link else BLACK)
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 10
    for m in range(MONTHS):
        ws.column_dimensions[get_column_letter(FIRST_COL + m)].width = 10
    ws.freeze_panes = "C5"
    return ws


build_model("Model_Base", "B")
build_model("Model_Growth", "C")
R = MODEL_ROWS
M12 = get_column_letter(FIRST_COL + 11)

# -------------------------------------------------------------- unit economics
U = wb.create_sheet("Unit_Economics")
U["A1"] = "Unit economics: one watchlist subscription, per month"
U["A1"].font = font(bold=True, size=13)
for i, h in enumerate(["Line", "Base", "Growth", "Unit"], 1):
    c = U.cell(row=3, column=i, value=h)
    c.font = font(bold=True)
    c.fill = HEADER_FILL
ue = [
    ("Price", "=" + a("price", "{s}"), "$ / month", USD_CENTS),
    ("Payment fees", "=%s*%s+%s" % (a("price", "{s}"), a("fee_pct", "{s}"), a("fee_fixed", "{s}")), "$ / month", USD_CENTS),
    ("LLM inference", "=%s*%s" % (a("decisions", "{s}"), a("llm", "{s}")), "$ / month", USD_CENTS),
    ("Laya inference", "=%s*%s" % (a("decisions", "{s}"), a("laya", "{s}")), "$ / month", USD_CENTS),
    ("Contribution margin", "={c}4-{c}5-{c}6-{c}7", "$ / month", USD_CENTS),
    ("Contribution margin, share of price", "=IF({c}4=0,0,{c}8/{c}4)", "%", PCT),
    ("Acquisition cost (CAC)", "=" + a("cac", "{s}"), "$ / customer", USD),
    ("CAC payback", '=IF({c}8<=0,"n/a",{c}10/{c}8)', "months", "0.0"),
    ("Customer lifetime value (LTV)", '=IF(%s=0,"n/a",{c}8/%s)' % (a("churn", "{s}"), a("churn", "{s}")), "$", USD),
    ("LTV to CAC", '=IF({c}10=0,"n/a",{c}12/{c}10)', "x", '0.0"x"'),
    ("LLM cost per decision at off-peak rates", "=%s/2" % a("llm", "{s}"), "$ / decision", USD_MILLS),
]
for i, (label, tmpl, unit, fmt) in enumerate(ue):
    r = 4 + i
    U.cell(row=r, column=1, value=label).font = font(bold=label in ("Contribution margin", "LTV to CAC"))
    for col, scen in ((2, "B"), (3, "C")):
        v = tmpl.format(s=scen, c=get_column_letter(col))
        cell = U.cell(row=r, column=col, value=v)
        cell.number_format = fmt
        cell.font = font(GREEN if v.startswith("=Assumptions!") and "*" not in v and "/" not in v else BLACK)
    U.cell(row=r, column=4, value=unit).font = font(size=9)
U["A16"] = "LTV = contribution margin / monthly churn. Payback = CAC / contribution margin."
U["A16"].font = font(italic=True, size=9)
U.column_dimensions["A"].width = 42
for col in "BC":
    U.column_dimensions[col].width = 13
U.column_dimensions["D"].width = 14

# ---------------------------------------------------------------------- summary
S = wb.create_sheet("Summary", 0)
S["A1"] = "Trading swarm: business plan summary"
S["A1"].font = font(bold=True, size=13)
S["A2"] = "Two scenarios over 24 months. Change inputs on the Assumptions sheet; everything here recalculates."
S["A2"].font = font(italic=True, size=9)
for i, h in enumerate(["Metric", "Base", "Growth", "Unit"], 1):
    c = S.cell(row=4, column=i, value=h)
    c.font = font(bold=True)
    c.fill = HEADER_FILL
summary = [
    ("Customers, month 12", "={m}!%s%d" % (M12, R["cust_end"]), "customers", COUNT),
    ("Customers, month 24", "={m}!%s%d" % (LAST, R["cust_end"]), "customers", COUNT),
    ("MRR, month 12", "={m}!%s%d" % (M12, R["revenue"]), "$ / month", USD),
    ("MRR, month 24", "={m}!%s%d" % (LAST, R["revenue"]), "$ / month", USD),
    ("Gross margin, month 24", "={m}!%s%d" % (LAST, R["gm"]), "%", PCT),
    ("Revenue, months 1-24", "=SUM({m}!C%d:%s%d)" % (R["revenue"], LAST, R["revenue"]), "$", USD),
    ("Operating profit, months 1-24", "=SUM({m}!C%d:%s%d)" % (R["op"], LAST, R["op"]), "$", USD),
    ("Lowest cash balance", "=MIN({m}!C%d:%s%d)" % (R["cash_end"], LAST, R["cash_end"]), "$", USD),
    ("Cash, end of month 24", "={m}!%s%d" % (LAST, R["cash_end"]), "$", USD),
    ("Runway: first month cash goes below zero",
     '=IF(SUM({m}!C%d:%s%d)=0,"none within 24 months",MATCH(1,{m}!C%d:%s%d,0))' % (R["flag_cash"], LAST, R["flag_cash"], R["flag_cash"], LAST, R["flag_cash"]),
     "month", INT),
    ("Break-even: first month of operating profit",
     '=IF(SUM({m}!C%d:%s%d)=0,"not within 24 months",MATCH(1,{m}!C%d:%s%d,0))' % (R["flag_profit"], LAST, R["flag_profit"], R["flag_profit"], LAST, R["flag_profit"]),
     "month", INT),
    ("Funding needed to never go below zero", "=MAX(0,-MIN({m}!C%d:%s%d))" % (R["cash_end"], LAST, R["cash_end"]), "$", USD),
    ("Contribution margin per watchlist", "=Unit_Economics!{u}8", "$ / month", USD_CENTS),
    ("LTV to CAC", "=Unit_Economics!{u}13", "x", '0.0"x"'),
]
for i, (label, tmpl, unit, fmt) in enumerate(summary):
    r = 5 + i
    S.cell(row=r, column=1, value=label).font = font()
    for col, m, u in ((2, "Model_Base", "B"), (3, "Model_Growth", "C")):
        cell = S.cell(row=r, column=col, value=tmpl.format(m=m, u=u))
        cell.number_format = fmt
        cell.font = font(GREEN if tmpl.startswith("={m}!") or tmpl.startswith("=Unit") else BLACK)
    S.cell(row=r, column=4, value=unit).font = font(size=9)
notes = [
    "How to read this:",
    "- Variable cost per watchlist is driven by measured LLM spend ($0.038 per decision, an upper bound at peak prices).",
    "- The swarm's validation result (1 of 577 setups approved) is a product property: it refuses trades without evidence.",
    "  The plan does not assume the swarm makes money for customers.",
    "- Customer numbers, price, churn, CAC, team cost and starting cash are owner estimates (yellow on Assumptions).",
    "- No fees or taxes beyond card processing are modeled.",
]
for i, t in enumerate(notes):
    c = S.cell(row=21 + i, column=1, value=t)
    c.font = font(bold=i == 0, size=9)
S.column_dimensions["A"].width = 46
for col in "BC":
    S.column_dimensions[col].width = 22
S.column_dimensions["D"].width = 12

for ws in wb.worksheets:
    style_all(ws)
wb.save(OUT)
print("saved", OUT, "| model rows:", {k: R[k] for k in ("cust_end", "revenue", "op", "cash_end", "flag_cash", "flag_profit")})
